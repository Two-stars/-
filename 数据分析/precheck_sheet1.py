# -*- coding: utf-8 -*-
"""Precheck and front analysis for logging fluid-identification data.

Example:
    python precheck_sheet1.py --input data.xlsx --sheet Sheet1 --output outputs_precheck
"""

from __future__ import annotations

import argparse
import json
import math
import os
import sys
import warnings
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence, Tuple

import matplotlib

matplotlib.use("Agg")

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd

# Older scikit-learn versions still reference removed NumPy aliases. Restoring
# them locally keeps the analysis runnable on the bundled Anaconda environment.
np.int = int  # type: ignore[attr-defined]
np.bool = bool  # type: ignore[attr-defined]
np.float = float  # type: ignore[attr-defined]

os.environ.setdefault("PYTHONWARNINGS", "ignore")
warnings.filterwarnings("ignore", category=DeprecationWarning)
warnings.filterwarnings("ignore", category=PendingDeprecationWarning)
warnings.filterwarnings("ignore", category=FutureWarning)
if hasattr(np, "VisibleDeprecationWarning"):
    warnings.filterwarnings("ignore", category=np.VisibleDeprecationWarning)

try:
    import openpyxl
except ImportError as exc:  # pragma: no cover - handled at runtime
    raise SystemExit("openpyxl is required to read the Excel workbook.") from exc

from sklearn.ensemble import RandomForestClassifier
from sklearn.feature_selection import mutual_info_classif
from sklearn.preprocessing import LabelEncoder

try:
    from sklearn.impute import SimpleImputer
except ImportError:  # Older scikit-learn compatibility.
    class SimpleImputer:  # type: ignore[no-redef]
        def __init__(self, strategy: str = "median") -> None:
            if strategy != "median":
                raise ValueError("Fallback SimpleImputer only supports median strategy.")
            self.statistics_: Optional[pd.Series] = None

        def fit_transform(self, x: pd.DataFrame) -> np.ndarray:
            self.statistics_ = x.median(axis=0).fillna(0)
            return x.fillna(self.statistics_).values


TARGET_CLASSES = ["气层", "含气层", "含气水层", "气水同层", "水层", "干层"]
LABEL_COL = "试油结论"
GROUP_COL = "井名"
DEPTH_COL = "深度"

CORE_FEATURES = ["GR", "SP", "AC", "CAL", "DEN", "CNL", "PE", "RD", "RS", "POR", "SW", "PERM"]
CANDIDATE_FEATURES = ["AC1", "DEN1", "CNL1", "U", "log（RD）", "log（RS）", "m2r3", "m2r9", "∆ Φ1"]
FIXED_EXCLUDED = [GROUP_COL, DEPTH_COL, LABEL_COL, "层位", "备注", "m2r6", "∆ Φ2", "∆ Φ3"]
EXPECTED_COLUMNS = [
    GROUP_COL,
    DEPTH_COL,
    "GR",
    "SP",
    "AC",
    "AC1",
    "CAL",
    "U",
    "DEN",
    "DEN1",
    "CNL",
    "CNL1",
    "PE",
    "RD",
    "RS",
    "log（RD）",
    "log（RS）",
    "m2r3",
    "m2r6",
    "m2r9",
    "POR",
    "SW",
    "PERM",
    LABEL_COL,
    "层位",
    "∆ Φ1",
    "∆ Φ2",
    "∆ Φ3",
    "备注",
]
DERIVED_FEATURES = [
    "RD_RS_ratio",
    "RD_RS_diff",
    "logRD_logRS_diff",
    "m2r9_m2r3_diff",
    "m2r9_m2r3_ratio",
    "CNL_DEN_diff",
    "AC_DEN_diff",
]
COARSE_LABEL_MAPPING = {
    "气层": "气层优势类",
    "含气层": "气层优势类",
    "含气水层": "气水风险类",
    "气水同层": "气水风险类",
    "水层": "非气层类",
    "干层": "非气层类",
}


def configure_stdout() -> None:
    """Avoid Windows console encoding failures when printing Chinese paths."""
    try:
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    except Exception:
        pass


def configure_matplotlib_fonts() -> str:
    """Use a common Chinese font if one is available; plotting still works otherwise."""
    candidates = [
        "Microsoft YaHei",
        "SimHei",
        "SimSun",
        "Noto Sans CJK SC",
        "Source Han Sans SC",
        "Arial Unicode MS",
        "DejaVu Sans",
    ]
    available = {f.name for f in matplotlib.font_manager.fontManager.ttflist}
    selected = next((font for font in candidates if font in available), "DejaVu Sans")
    plt.rcParams["font.sans-serif"] = [selected]
    plt.rcParams["axes.unicode_minus"] = False
    return selected


def ensure_dir(path: Path) -> None:
    path.mkdir(parents=True, exist_ok=True)


def read_excel_with_openpyxl(path: Path, sheet_name: str) -> Tuple[pd.DataFrame, pd.DataFrame]:
    """Read both cached values and formula/raw values from an Excel sheet."""
    if not path.exists():
        raise FileNotFoundError(f"Input file not found: {path}")

    def workbook_to_df(data_only: bool) -> pd.DataFrame:
        wb = openpyxl.load_workbook(str(path), data_only=data_only, read_only=True)
        if sheet_name not in wb.sheetnames:
            wb.close()
            raise ValueError(f"Sheet '{sheet_name}' not found. Available sheets: {wb.sheetnames}")

        ws = wb[sheet_name]
        rows = ws.iter_rows()
        try:
            header_cells = next(rows)
        except StopIteration:
            wb.close()
            return pd.DataFrame()

        headers = [str(cell.value).strip() if cell.value is not None else f"Unnamed_{i+1}" for i, cell in enumerate(header_cells)]
        records = []
        for row in rows:
            values = [cell.value for cell in row[: len(headers)]]
            if len(values) < len(headers):
                values.extend([None] * (len(headers) - len(values)))
            records.append(values)
        wb.close()
        return pd.DataFrame(records, columns=headers)

    return workbook_to_df(True), workbook_to_df(False)


def is_minus_9999(value: object) -> bool:
    if value is None:
        return False
    if isinstance(value, str):
        text = value.strip()
        if not text:
            return False
        try:
            return float(text) == -9999.0
        except ValueError:
            return False
    try:
        return float(value) == -9999.0
    except (TypeError, ValueError):
        return False


def is_ref_error(value: object) -> bool:
    return isinstance(value, str) and "#REF!" in value.upper()


def is_empty_like(value: object) -> bool:
    return value is None or (isinstance(value, str) and value.strip() == "")


def normalize_missing(df: pd.DataFrame) -> pd.DataFrame:
    cleaned = df.copy()
    for col in cleaned.columns:
        cleaned[col] = cleaned[col].map(lambda x: np.nan if is_empty_like(x) or is_minus_9999(x) or is_ref_error(x) else x)
    return cleaned


def numeric_series(df: pd.DataFrame, column: str) -> pd.Series:
    if column not in df.columns:
        return pd.Series(dtype=float)
    return pd.to_numeric(df[column], errors="coerce")


def add_or_repair_log_columns(df: pd.DataFrame) -> Tuple[pd.DataFrame, List[str]]:
    cleaned = df.copy()
    notes: List[str] = []
    rd = numeric_series(cleaned, "RD")
    rs = numeric_series(cleaned, "RS")

    for source, log_col in [("RD", "log（RD）"), ("RS", "log（RS）")]:
        source_values = numeric_series(cleaned, source)
        computed = pd.Series(np.where(source_values > 0, np.log10(source_values), np.nan), index=cleaned.index)
        if log_col in cleaned.columns:
            current = pd.to_numeric(cleaned[log_col], errors="coerce")
            missing_before = int(current.isna().sum())
            cleaned[log_col] = current.fillna(computed)
            repaired = int((current.isna() & computed.notna()).sum())
            if repaired:
                notes.append(f"{log_col} 有 {repaired} 个缺失值由 {source}>0 时的 log10({source}) 补算。")
            elif missing_before:
                notes.append(f"{log_col} 仍有 {missing_before} 个缺失值无法由 {source} 补算。")
        else:
            cleaned[log_col] = computed
            notes.append(f"{log_col} 原字段缺失，已在 {source}>0 时由 log10({source}) 计算生成。")

    return cleaned, notes


def safe_divide(numerator: pd.Series, denominator: pd.Series) -> pd.Series:
    denominator = denominator.replace(0, np.nan)
    with np.errstate(divide="ignore", invalid="ignore"):
        result = numerator / denominator
    return result.replace([np.inf, -np.inf], np.nan)


def build_derived_features(df: pd.DataFrame) -> Tuple[pd.DataFrame, Dict[str, str]]:
    result = df.copy()
    status: Dict[str, str] = {}

    rd = numeric_series(result, "RD")
    rs = numeric_series(result, "RS")
    log_rd = numeric_series(result, "log（RD）")
    log_rs = numeric_series(result, "log（RS）")
    m2r9 = numeric_series(result, "m2r9")
    m2r3 = numeric_series(result, "m2r3")
    cnl = numeric_series(result, "CNL")
    den = numeric_series(result, "DEN")
    ac = numeric_series(result, "AC")

    calculations = {
        "RD_RS_ratio": safe_divide(rd, rs),
        "RD_RS_diff": rd - rs,
        "logRD_logRS_diff": log_rd - log_rs,
        "m2r9_m2r3_diff": m2r9 - m2r3,
        "m2r9_m2r3_ratio": safe_divide(m2r9, m2r3),
        "CNL_DEN_diff": cnl - den,
        "AC_DEN_diff": ac - den,
    }
    for col, values in calculations.items():
        result[col] = values
        valid = int(values.notna().sum())
        status[col] = f"可计算 {valid} 条，占 {valid / max(len(result), 1):.2%}。"
        if valid == 0:
            status[col] += " 上游字段缺失或分母为 0，无法计算。"
    return result, status


def summarize_quality(
    raw_df: pd.DataFrame,
    value_df: pd.DataFrame,
    cleaned_df: pd.DataFrame,
    selected_core: Sequence[str],
    selected_candidate: Sequence[str],
    excluded_features: Sequence[str],
) -> pd.DataFrame:
    rows = []
    all_columns = list(dict.fromkeys(list(raw_df.columns) + EXPECTED_COLUMNS + DERIVED_FEATURES))
    for col in all_columns:
        if col not in cleaned_df.columns:
            rows.append(
                {
                    "字段名": col,
                    "数据类型": "字段缺失",
                    "非空数量": 0,
                    "缺失数量": len(cleaned_df),
                    "缺失比例": 1.0,
                    "-9999数量": 0,
                    "#REF!数量": 0,
                    "非法字符数量": 0,
                    "最小值": np.nan,
                    "最大值": np.nan,
                    "均值": np.nan,
                    "中位数": np.nan,
                    "标准差": np.nan,
                    "是否建议进入模型": "否",
                    "原因": "字段不存在。",
                }
            )
            continue

        raw_s = raw_df[col] if col in raw_df.columns else pd.Series([np.nan] * len(cleaned_df))
        value_s = value_df[col] if col in value_df.columns else pd.Series([np.nan] * len(cleaned_df))
        clean_s = cleaned_df[col]
        numeric = pd.to_numeric(clean_s, errors="coerce")
        numeric_like_nonmissing = clean_s.notna()
        illegal_count = int((numeric_like_nonmissing & numeric.isna() & ~clean_s.astype(str).str.strip().isin(TARGET_CLASSES)).sum())

        missing_count = int(clean_s.isna().sum())
        missing_ratio = missing_count / max(len(clean_s), 1)
        minus_count = int(raw_s.map(is_minus_9999).sum())
        ref_count = int((raw_s.map(is_ref_error) | value_s.map(is_ref_error)).sum())

        if col in selected_core:
            recommend = "是"
            reason = "核心测井字段，缺失比例较低，建议进入第一版模型。"
        elif col in selected_candidate:
            recommend = "候选"
            reason = "候选字段，数据质量可接受，可进入第一版模型或用于对比实验。"
        elif col in excluded_features:
            recommend = "否"
            if col == GROUP_COL:
                reason = "分组字段，仅用于按井验证，不作为模型输入。"
            elif col == DEPTH_COL:
                reason = "深度字段第一版不输入，仅用于排序和后续深度窗口特征。"
            elif col == LABEL_COL:
                reason = "标签字段，不能作为输入。"
            elif col == "层位":
                reason = "层位第一版暂不输入，需先评估层位分布和跨井泛化风险。"
            elif col == "备注":
                reason = "备注可能包含试油方式等强标签信息，存在标签泄漏风险。"
            elif col == "m2r6":
                reason = "按方案第一版排除；该字段常见 -9999 或缺失风险。"
            elif col in ["∆ Φ2", "∆ Φ3"]:
                reason = "按方案第一版排除；公式字段存在错误或泄漏风险。"
            else:
                reason = "按方案暂不作为第一版输入。"
        elif col in CORE_FEATURES or col in CANDIDATE_FEATURES:
            recommend = "否"
            reason = "字段缺失比例较高或有效数不足，暂不进入第一版模型。"
        elif col in DERIVED_FEATURES:
            valid = int(numeric.notna().sum())
            recommend = "是" if valid > 0 else "否"
            reason = "低对比响应增强特征。" if valid > 0 else "低对比响应增强特征无法计算。"
        else:
            recommend = "否"
            reason = "不在第一版输入清单内。"

        rows.append(
            {
                "字段名": col,
                "数据类型": str(clean_s.dtype),
                "非空数量": int(clean_s.notna().sum()),
                "缺失数量": missing_count,
                "缺失比例": missing_ratio,
                "-9999数量": minus_count,
                "#REF!数量": ref_count,
                "非法字符数量": illegal_count,
                "最小值": numeric.min(skipna=True),
                "最大值": numeric.max(skipna=True),
                "均值": numeric.mean(skipna=True),
                "中位数": numeric.median(skipna=True),
                "标准差": numeric.std(skipna=True),
                "是否建议进入模型": recommend,
                "原因": reason,
            }
        )
    return pd.DataFrame(rows)


def choose_features(cleaned_df: pd.DataFrame) -> Tuple[List[str], List[str], List[str]]:
    selected_core: List[str] = []
    selected_candidate: List[str] = []
    quality_excluded: List[str] = []

    def usable_numeric(col: str, max_missing_ratio: float) -> bool:
        if col not in cleaned_df.columns:
            return False
        numeric = pd.to_numeric(cleaned_df[col], errors="coerce")
        non_null = int(numeric.notna().sum())
        missing_ratio = 1 - non_null / max(len(cleaned_df), 1)
        return non_null > 0 and missing_ratio <= max_missing_ratio

    for col in CORE_FEATURES:
        if usable_numeric(col, 0.35):
            selected_core.append(col)
        else:
            quality_excluded.append(col)

    for col in CANDIDATE_FEATURES:
        if usable_numeric(col, 0.35):
            selected_candidate.append(col)
        else:
            quality_excluded.append(col)

    return selected_core, selected_candidate, quality_excluded


def save_class_distribution(df: pd.DataFrame, output_dir: Path, font_name: str) -> pd.DataFrame:
    label = df[LABEL_COL].astype(str).str.strip()
    counts = label.value_counts(dropna=False).reindex(TARGET_CLASSES).fillna(0).astype(int)
    distribution = pd.DataFrame({"类别": counts.index, "样本数": counts.values})
    distribution["比例"] = distribution["样本数"] / max(distribution["样本数"].sum(), 1)
    distribution.to_csv(output_dir / "class_distribution.csv", index=False, encoding="utf-8-sig")

    plt.figure(figsize=(8, 4.8), dpi=160)
    bars = plt.bar(distribution["类别"], distribution["样本数"], color="#4E79A7")
    plt.title("试油结论六分类分布")
    plt.xlabel("类别")
    plt.ylabel("样本数")
    plt.xticks(rotation=25, ha="right")
    for bar, ratio in zip(bars, distribution["比例"]):
        plt.text(bar.get_x() + bar.get_width() / 2, bar.get_height(), f"{ratio:.1%}", ha="center", va="bottom", fontsize=9)
    plt.tight_layout()
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        plt.savefig(output_dir / "class_distribution.png", bbox_inches="tight")
    plt.close()
    return distribution


def save_well_distribution(df: pd.DataFrame, output_dir: Path) -> Tuple[pd.DataFrame, pd.DataFrame]:
    if GROUP_COL not in df.columns:
        empty = pd.DataFrame()
        empty.to_csv(output_dir / "well_sample_count.csv", index=False, encoding="utf-8-sig")
        empty.to_csv(output_dir / "well_class_crosstab.csv", index=False, encoding="utf-8-sig")
        return empty, empty

    well_counts = df[GROUP_COL].value_counts(dropna=False).rename_axis(GROUP_COL).reset_index(name="样本数")
    well_counts.to_csv(output_dir / "well_sample_count.csv", index=False, encoding="utf-8-sig")

    if LABEL_COL in df.columns:
        crosstab = pd.crosstab(df[GROUP_COL], df[LABEL_COL]).reindex(columns=TARGET_CLASSES, fill_value=0)
        crosstab.to_csv(output_dir / "well_class_crosstab.csv", encoding="utf-8-sig")
    else:
        crosstab = pd.DataFrame()
        crosstab.to_csv(output_dir / "well_class_crosstab.csv", encoding="utf-8-sig")
    return well_counts, crosstab


def save_feature_descriptive_stats(df: pd.DataFrame, features: Sequence[str], output_dir: Path) -> pd.DataFrame:
    rows = []
    for col in features:
        if col not in df.columns:
            rows.append({"字段名": col, "非空数量": 0, "缺失数量": len(df), "缺失比例": 1.0})
            continue
        numeric = pd.to_numeric(df[col], errors="coerce")
        desc = numeric.describe(percentiles=[0.25, 0.5, 0.75])
        rows.append(
            {
                "字段名": col,
                "非空数量": int(numeric.notna().sum()),
                "缺失数量": int(numeric.isna().sum()),
                "缺失比例": float(numeric.isna().mean()),
                "最小值": desc.get("min", np.nan),
                "25%": desc.get("25%", np.nan),
                "中位数": desc.get("50%", np.nan),
                "75%": desc.get("75%", np.nan),
                "最大值": desc.get("max", np.nan),
                "均值": desc.get("mean", np.nan),
                "标准差": desc.get("std", np.nan),
            }
        )
    stats = pd.DataFrame(rows)
    stats.to_csv(output_dir / "feature_descriptive_stats.csv", index=False, encoding="utf-8-sig")
    return stats


def save_boxplots(df: pd.DataFrame, features: Sequence[str], output_dir: Path) -> List[str]:
    box_dir = output_dir / "boxplots_key_features"
    ensure_dir(box_dir)
    saved: List[str] = []
    if LABEL_COL not in df.columns:
        return saved

    labels = [cls for cls in TARGET_CLASSES if cls in set(df[LABEL_COL].dropna().astype(str))]
    for col in features:
        if col not in df.columns:
            continue
        numeric = pd.to_numeric(df[col], errors="coerce")
        groups = [numeric[df[LABEL_COL].astype(str) == cls].dropna().values for cls in labels]
        if not labels or all(len(g) == 0 for g in groups):
            continue

        plt.figure(figsize=(8.5, 4.8), dpi=160)
        plt.boxplot(groups, labels=labels, showfliers=False, patch_artist=True)
        plt.title(f"{col} 按试油结论分布")
        plt.xlabel("试油结论")
        plt.ylabel(col)
        plt.xticks(rotation=25, ha="right")
        plt.grid(axis="y", linestyle="--", alpha=0.3)
        plt.tight_layout()
        path = box_dir / f"{sanitize_filename(col)}.png"
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            plt.savefig(path, bbox_inches="tight")
        plt.close()
        saved.append(str(path.name))
    return saved


def sanitize_filename(name: str) -> str:
    return "".join(ch if ch not in '\\/:*?"<>|' else "_" for ch in name)


def save_correlation(df: pd.DataFrame, output_dir: Path) -> pd.DataFrame:
    numeric_df = pd.DataFrame()
    for col in df.columns:
        converted = pd.to_numeric(df[col], errors="coerce")
        if converted.notna().sum() > 0:
            numeric_df[col] = converted

    corr = numeric_df.corr(method="pearson")
    corr.to_csv(output_dir / "correlation_matrix.csv", encoding="utf-8-sig")

    if not corr.empty:
        fig_size = max(8, min(18, 0.45 * len(corr.columns) + 4))
        plt.figure(figsize=(fig_size, fig_size), dpi=160)
        im = plt.imshow(corr, cmap="coolwarm", vmin=-1, vmax=1)
        plt.colorbar(im, fraction=0.046, pad=0.04)
        plt.xticks(range(len(corr.columns)), corr.columns, rotation=60, ha="right", fontsize=8)
        plt.yticks(range(len(corr.index)), corr.index, fontsize=8)
        plt.title("Pearson 相关性矩阵")
        plt.tight_layout()
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            plt.savefig(output_dir / "correlation_matrix.png", bbox_inches="tight")
        plt.close()
    return corr


def prepare_feature_matrix(df: pd.DataFrame, features: Sequence[str]) -> Tuple[pd.DataFrame, List[str]]:
    usable = []
    x = pd.DataFrame(index=df.index)
    for col in features:
        if col not in df.columns:
            continue
        numeric = pd.to_numeric(df[col], errors="coerce")
        if numeric.notna().sum() > 0:
            x[col] = numeric
            usable.append(col)
    return x, usable


def save_feature_importance(df: pd.DataFrame, features: Sequence[str], output_dir: Path) -> Tuple[pd.DataFrame, pd.DataFrame]:
    x, usable = prepare_feature_matrix(df, features)
    if LABEL_COL not in df.columns or x.empty:
        rf = pd.DataFrame(columns=["feature", "importance"])
        mi = pd.DataFrame(columns=["feature", "mutual_info"])
        rf.to_csv(output_dir / "feature_importance_rf.csv", index=False, encoding="utf-8-sig")
        mi.to_csv(output_dir / "feature_importance_mi.csv", index=False, encoding="utf-8-sig")
        return rf, mi

    y_raw = df[LABEL_COL].astype(str).str.strip()
    mask = y_raw.isin(TARGET_CLASSES)
    x = x.loc[mask]
    y_raw = y_raw.loc[mask]
    if x.empty or y_raw.nunique() < 2:
        rf = pd.DataFrame(columns=["feature", "importance"])
        mi = pd.DataFrame(columns=["feature", "mutual_info"])
        rf.to_csv(output_dir / "feature_importance_rf.csv", index=False, encoding="utf-8-sig")
        mi.to_csv(output_dir / "feature_importance_mi.csv", index=False, encoding="utf-8-sig")
        return rf, mi

    imputer = SimpleImputer(strategy="median")
    x_imputed = pd.DataFrame(imputer.fit_transform(x), columns=x.columns, index=x.index)
    y = LabelEncoder().fit_transform(y_raw)

    rf_model = RandomForestClassifier(
        n_estimators=500,
        random_state=42,
        class_weight="balanced_subsample",
        n_jobs=1,
        min_samples_leaf=2,
    )
    rf_model.fit(x_imputed, y)
    rf = (
        pd.DataFrame({"feature": x_imputed.columns, "importance": rf_model.feature_importances_})
        .sort_values("importance", ascending=False)
        .reset_index(drop=True)
    )
    rf.to_csv(output_dir / "feature_importance_rf.csv", index=False, encoding="utf-8-sig")
    save_importance_plot(rf, "importance", "RandomForest 特征重要性", output_dir / "feature_importance_rf.png")

    mi_scores = mutual_info_classif(x_imputed, y, random_state=42, discrete_features=False)
    mi = pd.DataFrame({"feature": x_imputed.columns, "mutual_info": mi_scores}).sort_values("mutual_info", ascending=False).reset_index(drop=True)
    mi.to_csv(output_dir / "feature_importance_mi.csv", index=False, encoding="utf-8-sig")
    save_importance_plot(mi, "mutual_info", "Mutual Information 特征重要性", output_dir / "feature_importance_mi.png")
    return rf, mi


def save_importance_plot(table: pd.DataFrame, score_col: str, title: str, path: Path, top_n: int = 20) -> None:
    if table.empty:
        return
    data = table.head(top_n).iloc[::-1]
    plt.figure(figsize=(8.5, max(4.8, 0.32 * len(data) + 1.5)), dpi=160)
    plt.barh(data["feature"], data[score_col], color="#59A14F")
    plt.title(title)
    plt.xlabel(score_col)
    plt.tight_layout()
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        plt.savefig(path, bbox_inches="tight")
    plt.close()


def build_modeling_data(df: pd.DataFrame, model_features: Sequence[str], output_dir: Path) -> pd.DataFrame:
    cols = [col for col in [GROUP_COL, DEPTH_COL] if col in df.columns]
    feature_cols = [col for col in model_features if col in df.columns]
    out = df[cols + feature_cols].copy()
    for col in feature_cols:
        out[col] = pd.to_numeric(out[col], errors="coerce")
        if out[col].notna().sum() > 0:
            out[col] = out[col].fillna(out[col].median())

    out["y_fine"] = df[LABEL_COL].astype(str).str.strip() if LABEL_COL in df.columns else np.nan
    out["y_coarse"] = out["y_fine"].map(COARSE_LABEL_MAPPING)
    out = out[out["y_fine"].isin(TARGET_CLASSES)].copy()
    out.to_csv(output_dir / "modeling_data_point.csv", index=False, encoding="utf-8-sig")
    return out


def save_feature_config(
    selected_core: Sequence[str],
    selected_candidate: Sequence[str],
    derived_present: Sequence[str],
    excluded_features: Sequence[str],
    output_dir: Path,
) -> Dict[str, object]:
    config = {
        "core_features": list(selected_core),
        "candidate_features": list(selected_candidate),
        "derived_features": list(derived_present),
        "excluded_features": list(dict.fromkeys(excluded_features)),
        "label_column": LABEL_COL,
        "group_column": GROUP_COL,
        "depth_column": DEPTH_COL,
        "fine_label_mapping": {name: idx for idx, name in enumerate(TARGET_CLASSES)},
        "coarse_label_mapping": COARSE_LABEL_MAPPING,
    }
    with open(output_dir / "feature_config.json", "w", encoding="utf-8") as f:
        json.dump(config, f, ensure_ascii=False, indent=2)
    return config


def top_corr_pairs(corr: pd.DataFrame, top_n: int = 8) -> pd.DataFrame:
    if corr.empty:
        return pd.DataFrame(columns=["feature_1", "feature_2", "corr"])
    pairs = []
    cols = list(corr.columns)
    for i, left in enumerate(cols):
        for right in cols[i + 1 :]:
            value = corr.loc[left, right]
            if pd.notna(value):
                pairs.append((left, right, value, abs(value)))
    return (
        pd.DataFrame(pairs, columns=["feature_1", "feature_2", "corr", "abs_corr"])
        .sort_values("abs_corr", ascending=False)
        .head(top_n)
        .drop(columns=["abs_corr"])
    )


def class_feature_summary(df: pd.DataFrame, features: Sequence[str]) -> pd.DataFrame:
    rows = []
    if LABEL_COL not in df.columns:
        return pd.DataFrame(rows)
    for col in features:
        if col not in df.columns:
            continue
        numeric = pd.to_numeric(df[col], errors="coerce")
        grouped = numeric.groupby(df[LABEL_COL]).median()
        if grouped.notna().sum() == 0:
            continue
        max_cls = grouped.idxmax()
        min_cls = grouped.idxmin()
        rows.append(
            {
                "字段": col,
                "最高中位数类别": max_cls,
                "最高中位数": grouped.max(),
                "最低中位数类别": min_cls,
                "最低中位数": grouped.min(),
                "中位数极差": grouped.max() - grouped.min(),
            }
        )
    return pd.DataFrame(rows).sort_values("中位数极差", ascending=False)


def render_report(
    output_dir: Path,
    input_path: Path,
    sheet_name: str,
    raw_shape: Tuple[int, int],
    class_dist: pd.DataFrame,
    well_counts: pd.DataFrame,
    quality: pd.DataFrame,
    selected_core: Sequence[str],
    selected_candidate: Sequence[str],
    derived_present: Sequence[str],
    excluded_features: Sequence[str],
    log_notes: Sequence[str],
    derived_status: Dict[str, str],
    corr: pd.DataFrame,
    rf: pd.DataFrame,
    mi: pd.DataFrame,
    modeling_data: pd.DataFrame,
) -> None:
    total = raw_shape[0]
    n_wells = int(well_counts.shape[0]) if not well_counts.empty else 0
    nonzero_classes = class_dist[class_dist["样本数"] > 0]
    class_count = int(nonzero_classes.shape[0])
    imbalance_ratio = (
        float(nonzero_classes["样本数"].max() / max(nonzero_classes["样本数"].min(), 1)) if not nonzero_classes.empty else math.nan
    )

    issue_mask = (
        (quality["缺失比例"] >= 0.10)
        | (quality["-9999数量"] > 0)
        | (quality["#REF!数量"] > 0)
        | (quality["字段名"].isin(excluded_features))
    )
    high_missing = quality.loc[issue_mask, ["字段名", "缺失比例", "-9999数量", "#REF!数量", "是否建议进入模型", "原因"]].copy()
    top_quality_issues = high_missing.head(12)
    corr_pairs = top_corr_pairs(corr)
    class_summary = class_feature_summary(modeling_data.rename(columns={"y_fine": LABEL_COL}), CORE_FEATURES)

    def md_table(df: pd.DataFrame, max_rows: int = 12) -> str:
        if df.empty:
            return "无。"
        safe = df.head(max_rows).copy()
        for col in safe.columns:
            if pd.api.types.is_float_dtype(safe[col]):
                safe[col] = safe[col].map(lambda x: "" if pd.isna(x) else f"{x:.4g}")
            else:
                safe[col] = safe[col].map(lambda x: "" if pd.isna(x) else str(x))

        headers = [str(col) for col in safe.columns]

        def clean_cell(value: object) -> str:
            return str(value).replace("|", "\\|").replace("\n", " ")

        lines = [
            "| " + " | ".join(clean_cell(h) for h in headers) + " |",
            "| " + " | ".join("---" for _ in headers) + " |",
        ]
        for _, row in safe.iterrows():
            lines.append("| " + " | ".join(clean_cell(row[col]) for col in safe.columns) + " |")
        return "\n".join(lines)

    class_lines = class_dist.copy()
    class_lines["比例"] = class_lines["比例"].map(lambda x: f"{x:.2%}")

    rf_top = rf.head(10)
    mi_top = mi.head(10)

    report = f"""# 测井流体识别数据预处理与前置分析报告

## 1. 数据基本情况

- 输入文件：`{input_path.name}`
- 工作表：`{sheet_name}`
- 样本总数：{total}
- 字段数：{raw_shape[1]}
- 井数：{n_wells}
- 有效类别数：{class_count}
- 第一版建模数据行数：{len(modeling_data)}

本次分析将 `-9999`、`#REF!`、空字符串统一视为缺失值。对数曲线 `log（RD）`、`log（RS）` 若缺失，则在 `RD`、`RS` 大于 0 时用 `log10` 补算。中位数填补只用于导出第一版建模表，正式训练时仍应在训练集内部拟合填补器。

## 2. 类别不平衡情况

{md_table(class_lines, 20)}

最大类/最小类样本数比例约为 **{imbalance_ratio:.2f}**。这说明六分类任务存在类别不平衡，后续评估应避免只看总体准确率，建议同时关注宏平均 F1、每类召回率，并考虑 class weight 或重采样策略。

## 3. 井分布结论

每口井样本数和每口井类别交叉表已分别输出到 `well_sample_count.csv` 和 `well_class_crosstab.csv`。由于同一井内相邻深度样本高度相关，后续模型验证建议采用按 `井名` 分组的 GroupKFold 或留井验证，避免同一口井同时出现在训练集和验证集导致泛化评价偏乐观。

## 4. 字段数据质量总结

完整字段质量表见 `data_quality_report.csv`。缺失比例较高或明确排除字段如下：

{md_table(top_quality_issues, 12)}

## 5. 建议输入模型的字段

第一版建议进入模型的核心原始字段：

`{ "、".join(selected_core) if selected_core else "无" }`

数据质量可接受的候选字段：

`{ "、".join(selected_candidate) if selected_candidate else "无" }`

建议加入的低对比响应增强特征：

`{ "、".join(derived_present) if derived_present else "无" }`

增强特征可计算情况：

{chr(10).join(f"- `{k}`：{v}" for k, v in derived_status.items())}

对数曲线处理说明：

{chr(10).join(f"- {note}" for note in log_notes) if log_notes else "- 未发现需要补算的对数曲线缺失。"}

## 6. 不建议输入模型的字段及原因

第一版暂不输入：

`{ "、".join(excluded_features) }`

主要原因：`井名` 仅作分组验证；`深度` 第一版仅作排序和后续窗口特征；`试油结论` 是标签；`层位` 需先评估层位分布和跨井泛化；`备注` 可能包含试油方式等强标签信息，存在标签泄漏风险；`m2r6`、`∆ Φ2`、`∆ Φ3` 按方案先排除，避免缺失、公式错误或派生泄漏风险。

## 7. 关键特征分布分析结论

关键字段按类别箱线图已输出到 `boxplots_key_features/`。按各类别中位数极差排序，差异较明显的字段如下：

{md_table(class_summary, 12)}

这些差异只能说明单变量分布有区分度，不能直接代表最终模型性能；低对比度气层识别更依赖多曲线组合响应。

## 8. 相关性分析结论

Pearson 相关性矩阵已输出到 `correlation_matrix.csv` 和 `correlation_matrix.png`。绝对相关性最高的特征对如下：

{md_table(corr_pairs, 8)}

高度相关字段在树模型中通常可保留，但在线性模型或解释性分析中应注意共线性。

## 9. 特征重要性分析结论

RandomForest 前 10 个特征：

{md_table(rf_top, 10)}

Mutual Information 前 10 个特征：

{md_table(mi_top, 10)}

该重要性分析仅作为前置探索，未采用按井分组验证，不能作为最终模型效果评价。

## 10. 第一版模型输入建议

建议第一版使用：

- 原始核心字段：`{ "、".join(selected_core) if selected_core else "无" }`
- 候选字段：`{ "、".join(selected_candidate) if selected_candidate else "无" }`
- 增强字段：`{ "、".join(derived_present) if derived_present else "无" }`
- 标签：`y_fine` 六分类，`y_coarse` 三分类粗标签
- 分组字段：`井名`

已导出 `modeling_data_point.csv`，包含 `井名`、`深度`、第一版输入字段、增强特征、`y_fine` 和 `y_coarse`。

## 11. 后续建模建议

1. 使用按井分组的训练/验证划分，优先采用 GroupKFold 或 Leave-One-Well-Out。
2. 在交叉验证管线内部完成缺失值填补、标准化和特征选择，避免数据泄漏。
3. 对六分类类别不平衡问题，比较 class weight、分层采样和代价敏感学习。
4. 同时训练六分类 `y_fine` 和三分类 `y_coarse`，评估粗分类是否更稳定。
5. 结合地质认识检查 `备注`、`层位` 等非曲线字段的泄漏风险，再决定是否进入后续实验。
"""

    with open(output_dir / "precheck_report.md", "w", encoding="utf-8") as f:
        f.write(report)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Precheck Sheet1 logging fluid-identification data.")
    parser.add_argument("--input", required=True, help="Path to the input Excel file.")
    parser.add_argument("--sheet", default="Sheet1", help="Sheet name to analyze.")
    parser.add_argument("--output", default="outputs_precheck", help="Output directory.")
    return parser.parse_args()


def main() -> None:
    configure_stdout()
    font_name = configure_matplotlib_fonts()
    args = parse_args()
    input_path = Path(args.input)
    output_dir = Path(args.output)
    ensure_dir(output_dir)

    print(f"Reading {input_path} / {args.sheet} ...")
    value_df, raw_df = read_excel_with_openpyxl(input_path, args.sheet)
    cleaned = normalize_missing(value_df)
    cleaned, log_notes = add_or_repair_log_columns(cleaned)
    cleaned, derived_status = build_derived_features(cleaned)

    selected_core, selected_candidate, quality_excluded = choose_features(cleaned)
    derived_present = [col for col in DERIVED_FEATURES if col in cleaned.columns and pd.to_numeric(cleaned[col], errors="coerce").notna().sum() > 0]
    excluded_features = list(dict.fromkeys(FIXED_EXCLUDED + quality_excluded))

    quality = summarize_quality(raw_df, value_df, cleaned, selected_core, selected_candidate, excluded_features)
    quality.to_csv(output_dir / "data_quality_report.csv", index=False, encoding="utf-8-sig")

    class_dist = save_class_distribution(cleaned, output_dir, font_name)
    well_counts, _ = save_well_distribution(cleaned, output_dir)
    save_feature_descriptive_stats(cleaned, CORE_FEATURES + CANDIDATE_FEATURES, output_dir)
    save_boxplots(cleaned, CORE_FEATURES, output_dir)
    corr = save_correlation(cleaned, output_dir)

    importance_features = list(dict.fromkeys(selected_core + selected_candidate + derived_present))
    rf, mi = save_feature_importance(cleaned, importance_features, output_dir)
    modeling_data = build_modeling_data(cleaned, importance_features, output_dir)
    save_feature_config(selected_core, selected_candidate, derived_present, excluded_features, output_dir)

    render_report(
        output_dir=output_dir,
        input_path=input_path,
        sheet_name=args.sheet,
        raw_shape=value_df.shape,
        class_dist=class_dist,
        well_counts=well_counts,
        quality=quality,
        selected_core=selected_core,
        selected_candidate=selected_candidate,
        derived_present=derived_present,
        excluded_features=excluded_features,
        log_notes=log_notes,
        derived_status=derived_status,
        corr=corr,
        rf=rf,
        mi=mi,
        modeling_data=modeling_data,
    )
    print(f"Done. Outputs saved to {output_dir.resolve()}")


if __name__ == "__main__":
    main()
