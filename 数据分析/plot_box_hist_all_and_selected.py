# -*- coding: utf-8 -*-
"""Draw publication-style boxplots and histograms for logging indicators.

Example:
    python plot_box_hist_all_and_selected.py --input outputs_precheck/modeling_data_point.csv --output outputs_precheck
"""

from __future__ import annotations

import argparse
import sys
import warnings
from pathlib import Path
from typing import Dict, List, Sequence, Tuple

import matplotlib

matplotlib.use("Agg")

import matplotlib.lines as mlines
import matplotlib.patches as mpatches
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd

if hasattr(np, "VisibleDeprecationWarning"):
    warnings.filterwarnings("ignore", category=np.VisibleDeprecationWarning)


TARGET_CLASSES = ["气层", "含气层", "含气水层", "气水同层", "水层", "干层"]
LABEL_CANDIDATES = ["试油结论", "y_fine"]

RAW_PRIORITY_FEATURES = [
    "深度",
    "GR",
    "SP",
    "AC",
    "AC1",
    "CAL",
    "U",
    "DEN",
    "DEN1",
    "CNL",
    "CNL1",
    "PE",
    "RD",
    "RS",
    "log（RD）",
    "log（RS）",
    "m2r3",
    "m2r6",
    "m2r9",
    "POR",
    "SW",
    "PERM",
    "∆ Φ1",
    "∆ Φ2",
    "∆ Φ3",
]

DERIVED_FEATURES = [
    "RD_RS_ratio",
    "RD_RS_diff",
    "logRD_logRS_diff",
    "m2r9_m2r3_diff",
    "m2r9_m2r3_ratio",
    "CNL_DEN_diff",
    "AC_DEN_diff",
]

SELECTED_V1_FEATURES = ["GR", "SP", "AC", "CAL", "DEN", "CNL", "PE", "RD", "SW", "RD_RS_ratio"]
NON_NUMERIC_FIELDS = {"井名", "试油结论", "层位", "备注", "y_fine", "y_coarse"}
MIN_VALID_COUNT = 3


def configure_stdout() -> None:
    try:
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    except Exception:
        pass


def configure_matplotlib_fonts() -> str:
    candidates = [
        "Microsoft YaHei",
        "SimHei",
        "SimSun",
        "Noto Sans CJK SC",
        "Source Han Sans SC",
        "Arial Unicode MS",
        "DejaVu Sans",
    ]
    available = {f.name for f in matplotlib.font_manager.fontManager.ttflist}
    selected = next((font for font in candidates if font in available), "DejaVu Sans")
    plt.rcParams["font.sans-serif"] = [selected]
    plt.rcParams["axes.unicode_minus"] = False
    return selected


def sanitize_filename(name: str) -> str:
    return "".join(ch if ch not in '\\/:*?"<>|' else "_" for ch in name)


def read_excel_fallback(path: Path, sheet_name: str) -> pd.DataFrame:
    """Read Excel with pandas first, then openpyxl for older local environments."""
    try:
        return pd.read_excel(path, sheet_name=sheet_name)
    except Exception:
        import openpyxl

        wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
        if sheet_name not in wb.sheetnames:
            wb.close()
            raise ValueError(f"Sheet '{sheet_name}' not found in {path}.")
        ws = wb[sheet_name]
        rows = ws.iter_rows()
        header = [str(cell.value).strip() if cell.value is not None else f"Unnamed_{i + 1}" for i, cell in enumerate(next(rows))]
        records = []
        for row in rows:
            values = [cell.value for cell in row[: len(header)]]
            if len(values) < len(header):
                values.extend([np.nan] * (len(header) - len(values)))
            records.append(values)
        wb.close()
        return pd.DataFrame(records, columns=header)


def load_data(input_path: Path, sheet_name: str) -> Tuple[pd.DataFrame, Path]:
    if input_path.exists():
        if input_path.suffix.lower() in [".csv", ".txt"]:
            return pd.read_csv(input_path), input_path
        if input_path.suffix.lower() in [".xlsx", ".xlsm", ".xls"]:
            return read_excel_fallback(input_path, sheet_name), input_path
        raise ValueError(f"Unsupported input type: {input_path.suffix}")

    fallback = Path("流体识别数据.xlsx")
    if fallback.exists():
        print(f"提示：未找到 {input_path}，改用原始 Excel：{fallback}")
        return read_excel_fallback(fallback, sheet_name), fallback
    raise FileNotFoundError(f"Input file not found: {input_path}; fallback Excel also not found: {fallback}")


def is_bad_value(value: object) -> bool:
    if value is None:
        return True
    if isinstance(value, str):
        text = value.strip()
        if text == "":
            return True
        if "#REF!" in text.upper():
            return True
        try:
            return float(text) == -9999.0
        except ValueError:
            return False
    try:
        return float(value) == -9999.0
    except (TypeError, ValueError):
        return False


def clean_bad_values(df: pd.DataFrame) -> pd.DataFrame:
    cleaned = df.copy()
    for col in cleaned.columns:
        cleaned[col] = cleaned[col].map(lambda x: np.nan if is_bad_value(x) else x)
    return cleaned


def ensure_ratio(df: pd.DataFrame) -> Tuple[pd.DataFrame, str]:
    result = df.copy()
    if "RD_RS_ratio" in result.columns:
        return result, "RD_RS_ratio 已存在，直接使用。"
    if {"RD", "RS"}.issubset(result.columns):
        rd = pd.to_numeric(result["RD"], errors="coerce")
        rs = pd.to_numeric(result["RS"], errors="coerce").replace(0, np.nan)
        result["RD_RS_ratio"] = (rd / rs).replace([np.inf, -np.inf], np.nan)
        return result, "RD_RS_ratio 不存在，已由 RD / RS 自动计算。"
    return result, "RD_RS_ratio 不存在，且 RD 或 RS 缺失，无法计算。"


def find_label_column(df: pd.DataFrame) -> str:
    for col in LABEL_CANDIDATES:
        if col in df.columns:
            return col
    return ""


def prepare_numeric_feature(df: pd.DataFrame, feature: str) -> pd.Series:
    return pd.to_numeric(df[feature], errors="coerce").replace([np.inf, -np.inf], np.nan)


def collect_all_features(df: pd.DataFrame) -> List[str]:
    ordered = []
    for feature in RAW_PRIORITY_FEATURES + DERIVED_FEATURES:
        if feature in df.columns and feature not in ordered:
            ordered.append(feature)

    for feature in df.columns:
        if feature in NON_NUMERIC_FIELDS or feature in ordered:
            continue
        numeric = pd.to_numeric(df[feature], errors="coerce")
        if numeric.notna().sum() >= MIN_VALID_COUNT:
            ordered.append(feature)
    return ordered


def validate_features(df: pd.DataFrame, features: Sequence[str]) -> Tuple[List[str], Dict[str, str], Dict[str, int]]:
    ok: List[str] = []
    skipped: Dict[str, str] = {}
    valid_counts: Dict[str, int] = {}
    for feature in features:
        if feature not in df.columns:
            skipped[feature] = "字段不存在"
            print(f"跳过 {feature}: 字段不存在")
            continue
        numeric = prepare_numeric_feature(df, feature)
        valid_count = int(numeric.notna().sum())
        valid_counts[feature] = valid_count
        if valid_count == 0:
            skipped[feature] = "全为空或无法转为数值"
            print(f"跳过 {feature}: 全为空或无法转为数值")
        elif valid_count < MIN_VALID_COUNT:
            skipped[feature] = f"有效样本太少，仅 {valid_count} 条"
            print(f"跳过 {feature}: 有效样本太少，仅 {valid_count} 条")
        else:
            ok.append(feature)
    return ok, skipped, valid_counts


def groups_for_boxplot(df: pd.DataFrame, label_col: str, feature: str) -> Tuple[List[str], List[np.ndarray]]:
    labels = []
    groups = []
    numeric = prepare_numeric_feature(df, feature)
    label_values = df[label_col].astype(str).str.strip()
    for cls in TARGET_CLASSES:
        values = numeric[label_values == cls].dropna().values
        if len(values) >= 1:
            labels.append(cls)
            groups.append(values)
    return labels, groups


def plot_boxplot(df: pd.DataFrame, label_col: str, feature: str, output_path: Path, seed: int = 42) -> bool:
    labels, groups = groups_for_boxplot(df, label_col, feature)
    if not groups:
        print(f"跳过 {feature} 箱线图: 没有可按类别绘制的有效数据")
        return False

    rng = np.random.default_rng(seed)
    fig, ax = plt.subplots(figsize=(12, 6), dpi=300)
    positions = np.arange(1, len(groups) + 1)

    ax.boxplot(
        groups,
        positions=positions,
        widths=0.58,
        whis=(0, 100),
        patch_artist=True,
        showmeans=True,
        showfliers=False,
        boxprops={"facecolor": "#CFEA8D", "edgecolor": "#1F77B4", "linewidth": 1.4},
        whiskerprops={"color": "#1F77B4", "linewidth": 1.2},
        capprops={"color": "#1F77B4", "linewidth": 1.2},
        medianprops={"color": "#1F77B4", "linewidth": 1.8},
        meanprops={
            "marker": "s",
            "markerfacecolor": "#E15759",
            "markeredgecolor": "#E15759",
            "markersize": 4.5,
        },
    )

    for pos, values in zip(positions, groups):
        jitter = rng.uniform(-0.16, 0.16, size=len(values))
        ax.scatter(
            np.full(len(values), pos) + jitter,
            values,
            s=12,
            facecolors="none",
            edgecolors="#7B3F98",
            linewidths=0.55,
            alpha=0.65,
            zorder=2,
        )

    ax.set_title(f"Box plot of {feature}", fontsize=14)
    ax.set_xlabel("试油结论")
    ax.set_ylabel(feature)
    ax.set_xticks(positions)
    ax.set_xticklabels(labels)
    ax.grid(axis="y", linestyle="--", alpha=0.28)

    legend_handles = [
        mpatches.Patch(facecolor="#CFEA8D", edgecolor="#1F77B4", label="25%–75%"),
        mlines.Line2D([], [], color="#1F77B4", linewidth=1.2, label="Max–Min"),
        mlines.Line2D([], [], color="#1F77B4", linewidth=1.8, label="Median line"),
        mlines.Line2D([], [], color="#E15759", marker="s", linestyle="None", markersize=5, label="Mean"),
    ]
    ax.legend(handles=legend_handles, loc="best", frameon=True, fontsize=9)

    fig.tight_layout()
    fig.savefig(output_path, bbox_inches="tight")
    plt.close(fig)
    return True


def plot_histogram(df: pd.DataFrame, feature: str, output_path: Path) -> bool:
    values = prepare_numeric_feature(df, feature).dropna()
    if len(values) < MIN_VALID_COUNT:
        print(f"跳过 {feature} 直方图: 有效样本太少，仅 {len(values)} 条")
        return False

    fig, ax = plt.subplots(figsize=(8, 6), dpi=300)
    ax.hist(values, bins=30, color="#1F77B4", edgecolor="black", linewidth=0.5, alpha=0.85)
    ax.set_title(f"Histogram of {feature}", fontsize=14)
    ax.set_xlabel(feature)
    ax.set_ylabel("Frequency")
    ax.grid(axis="y", linestyle="--", alpha=0.25)
    fig.tight_layout()
    fig.savefig(output_path, bbox_inches="tight")
    plt.close(fig)
    return True


def draw_feature_set(
    df: pd.DataFrame,
    label_col: str,
    features: Sequence[str],
    box_dir: Path,
    hist_dir: Path,
) -> Tuple[List[str], List[str]]:
    box_dir.mkdir(parents=True, exist_ok=True)
    hist_dir.mkdir(parents=True, exist_ok=True)
    box_success: List[str] = []
    hist_success: List[str] = []
    for feature in features:
        safe = sanitize_filename(feature)
        if plot_boxplot(df, label_col, feature, box_dir / f"boxplot_{safe}.png"):
            box_success.append(feature)
        if plot_histogram(df, feature, hist_dir / f"histogram_{safe}.png"):
            hist_success.append(feature)
    return box_success, hist_success


def write_summary(
    output_dir: Path,
    actual_input: Path,
    label_col: str,
    ratio_note: str,
    all_features: Sequence[str],
    selected_features: Sequence[str],
    skipped_all: Dict[str, str],
    skipped_selected: Dict[str, str],
    all_box_count: int,
    all_hist_count: int,
    selected_box_count: int,
    selected_hist_count: int,
    font_name: str,
) -> None:
    def bullet_list(items: Sequence[str]) -> str:
        return "\n".join(f"- `{item}`" for item in items) if items else "- 无"

    def skipped_table(skipped: Dict[str, str]) -> str:
        if not skipped:
            return "无。"
        lines = ["| 字段 | 跳过原因 |", "| --- | --- |"]
        for field, reason in skipped.items():
            lines.append(f"| `{field}` | {reason} |")
        return "\n".join(lines)

    report = f"""# 绘图结果汇总

## 数据来源

- 输入文件：`{actual_input}`
- 分组标签列：`{label_col}`
- 字体：`{font_name}`
- `RD_RS_ratio` 处理：{ratio_note}

## 成功绘图的所有指标

{bullet_list(all_features)}

## 成功绘图的第一版输入特征

{bullet_list(selected_features)}

## 跳过字段：所有指标

{skipped_table(skipped_all)}

## 跳过字段：第一版输入特征

{skipped_table(skipped_selected)}

## 输出目录

- 所有指标箱线图：`{output_dir / "plots_all" / "boxplots"}`
- 所有指标直方图：`{output_dir / "plots_all" / "histograms"}`
- 第一版输入特征箱线图：`{output_dir / "plots_selected_v1" / "boxplots"}`
- 第一版输入特征直方图：`{output_dir / "plots_selected_v1" / "histograms"}`

## 图片数量

- 所有指标箱线图：{all_box_count}
- 所有指标直方图：{all_hist_count}
- 第一版输入特征箱线图：{selected_box_count}
- 第一版输入特征直方图：{selected_hist_count}
"""
    (output_dir / "plot_summary.md").write_text(report, encoding="utf-8")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Draw boxplots and histograms for all and selected logging indicators.")
    parser.add_argument("--input", default="outputs_precheck/modeling_data_point.csv", help="Preferred input CSV or Excel path.")
    parser.add_argument("--output", default="outputs_precheck", help="Output directory.")
    parser.add_argument("--sheet", default="Sheet1", help="Excel sheet name used only for Excel input or fallback.")
    return parser.parse_args()


def main() -> None:
    configure_stdout()
    font_name = configure_matplotlib_fonts()
    args = parse_args()

    input_path = Path(args.input)
    output_dir = Path(args.output)
    output_dir.mkdir(parents=True, exist_ok=True)

    df, actual_input = load_data(input_path, args.sheet)
    df = clean_bad_values(df)
    df, ratio_note = ensure_ratio(df)

    label_col = find_label_column(df)
    if not label_col:
        raise ValueError("未找到箱线图分组标签列：试油结论 或 y_fine。")

    all_candidates = collect_all_features(df)
    all_features, skipped_all, _ = validate_features(df, all_candidates)

    selected_candidates = list(SELECTED_V1_FEATURES)
    selected_features, skipped_selected, _ = validate_features(df, selected_candidates)

    all_box, all_hist = draw_feature_set(
        df,
        label_col,
        all_features,
        output_dir / "plots_all" / "boxplots",
        output_dir / "plots_all" / "histograms",
    )
    selected_box, selected_hist = draw_feature_set(
        df,
        label_col,
        selected_features,
        output_dir / "plots_selected_v1" / "boxplots",
        output_dir / "plots_selected_v1" / "histograms",
    )

    # Include explicitly requested but absent priority fields in the summary.
    for feature in RAW_PRIORITY_FEATURES + DERIVED_FEATURES:
        if feature not in df.columns and feature not in skipped_all:
            skipped_all[feature] = "字段不存在"
    for feature in SELECTED_V1_FEATURES:
        if feature not in df.columns and feature not in skipped_selected:
            skipped_selected[feature] = "字段不存在"

    write_summary(
        output_dir=output_dir,
        actual_input=actual_input,
        label_col=label_col,
        ratio_note=ratio_note,
        all_features=all_box,
        selected_features=selected_box,
        skipped_all=skipped_all,
        skipped_selected=skipped_selected,
        all_box_count=len(all_box),
        all_hist_count=len(all_hist),
        selected_box_count=len(selected_box),
        selected_hist_count=len(selected_hist),
        font_name=font_name,
    )

    print(f"读取文件：{actual_input}")
    print(f"分组标签列：{label_col}")
    print(ratio_note)
    print(f"所有指标箱线图：{len(all_box)} 张")
    print(f"所有指标直方图：{len(all_hist)} 张")
    print(f"第一版输入特征箱线图：{len(selected_box)} 张")
    print(f"第一版输入特征直方图：{len(selected_hist)} 张")
    print(f"plot_summary：{(output_dir / 'plot_summary.md').resolve()}")
    if skipped_all or skipped_selected:
        print("存在跳过字段，详见 plot_summary.md。")


if __name__ == "__main__":
    main()
