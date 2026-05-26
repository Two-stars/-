# -*- coding: utf-8 -*-
"""Compute mutual-information rankings for logging fluid-identification features.

Example:
    python compute_mutual_information.py --input outputs_precheck/modeling_data_point.csv --output outputs_precheck/mi_analysis
"""

from __future__ import annotations

import argparse
import sys
import warnings
from pathlib import Path
from typing import Dict, List, Sequence, Tuple

import matplotlib

matplotlib.use("Agg")

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from sklearn.feature_selection import mutual_info_classif


LABEL_ORDER = ["气层", "含气层", "含气水层", "气水同层", "水层", "干层"]
LABEL_CANDIDATES = ["y_fine", "试油结论"]
MIN_VALID_COUNT = 30

CANDIDATE_FEATURES = [
    "GR",
    "SP",
    "AC",
    "AC1",
    "CAL",
    "U",
    "DEN",
    "DEN1",
    "CNL",
    "CNL1",
    "PE",
    "RD",
    "RS",
    "log（RD）",
    "log（RS）",
    "m2r3",
    "m2r9",
    "POR",
    "SW",
    "PERM",
    "∆ Φ1",
    "RD_RS_ratio",
    "RD_RS_diff",
    "logRD_logRS_diff",
    "m2r9_m2r3_diff",
    "m2r9_m2r3_ratio",
    "CNL_DEN_diff",
    "AC_DEN_diff",
]

SELECTED_V1_FEATURES = ["GR", "SP", "AC", "CAL", "DEN", "CNL", "PE", "RD", "SW", "RD_RS_ratio"]

EXCLUDED_FEATURES = {"井名", "深度", "试油结论", "层位", "备注", "y_fine", "y_coarse", "m2r6", "∆ Φ2", "∆ Φ3"}


def configure_stdout() -> None:
    try:
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    except Exception:
        pass


def configure_matplotlib_fonts() -> str:
    candidates = [
        "Microsoft YaHei",
        "SimHei",
        "SimSun",
        "Noto Sans CJK SC",
        "Source Han Sans SC",
        "Arial Unicode MS",
        "DejaVu Sans",
    ]
    available = {f.name for f in matplotlib.font_manager.fontManager.ttflist}
    selected = next((font for font in candidates if font in available), "DejaVu Sans")
    plt.rcParams["font.sans-serif"] = [selected]
    plt.rcParams["axes.unicode_minus"] = False
    return selected


def read_excel_fallback(path: Path, sheet_name: str) -> pd.DataFrame:
    """Read Excel with pandas first, then fall back to openpyxl for older envs."""
    try:
        return pd.read_excel(path, sheet_name=sheet_name)
    except Exception:
        import openpyxl

        wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
        if sheet_name not in wb.sheetnames:
            wb.close()
            raise ValueError(f"Sheet '{sheet_name}' not found in {path}.")
        ws = wb[sheet_name]
        rows = ws.iter_rows()
        header = [str(cell.value).strip() if cell.value is not None else f"Unnamed_{i + 1}" for i, cell in enumerate(next(rows))]
        records = []
        for row in rows:
            values = [cell.value for cell in row[: len(header)]]
            if len(values) < len(header):
                values.extend([np.nan] * (len(header) - len(values)))
            records.append(values)
        wb.close()
        return pd.DataFrame(records, columns=header)


def load_data(input_path: Path, sheet_name: str) -> Tuple[pd.DataFrame, Path]:
    if input_path.exists():
        if input_path.suffix.lower() in [".csv", ".txt"]:
            return pd.read_csv(input_path), input_path
        if input_path.suffix.lower() in [".xlsx", ".xlsm", ".xls"]:
            return read_excel_fallback(input_path, sheet_name), input_path
        raise ValueError(f"Unsupported input type: {input_path.suffix}")

    fallback = Path("流体识别数据.xlsx")
    if fallback.exists():
        print(f"提示：未找到 {input_path}，改用原始 Excel：{fallback}")
        return read_excel_fallback(fallback, sheet_name), fallback
    raise FileNotFoundError(f"Input file not found: {input_path}; fallback Excel also not found: {fallback}")


def is_bad_value(value: object) -> bool:
    if value is None:
        return True
    if isinstance(value, str):
        text = value.strip()
        if text == "":
            return True
        if "#REF!" in text.upper():
            return True
        try:
            return float(text) == -9999.0
        except ValueError:
            return False
    try:
        return float(value) == -9999.0
    except (TypeError, ValueError):
        return False


def clean_bad_values(df: pd.DataFrame) -> pd.DataFrame:
    cleaned = df.copy()
    for col in cleaned.columns:
        cleaned[col] = cleaned[col].map(lambda x: np.nan if is_bad_value(x) else x)
    return cleaned


def ensure_ratio(df: pd.DataFrame) -> Tuple[pd.DataFrame, str]:
    result = df.copy()
    if "RD_RS_ratio" in result.columns:
        return result, "RD_RS_ratio 已存在，直接使用。"
    if {"RD", "RS"}.issubset(result.columns):
        rd = pd.to_numeric(result["RD"], errors="coerce")
        rs = pd.to_numeric(result["RS"], errors="coerce").replace(0, np.nan)
        result["RD_RS_ratio"] = (rd / rs).replace([np.inf, -np.inf], np.nan)
        return result, "RD_RS_ratio 不存在，已由 RD / RS 自动计算。"
    return result, "RD_RS_ratio 不存在，且 RD 或 RS 缺失，无法计算。"


def find_label_column(df: pd.DataFrame) -> str:
    for col in LABEL_CANDIDATES:
        if col in df.columns:
            return col
    return ""


def encode_label(df: pd.DataFrame, label_col: str) -> Tuple[pd.Series, pd.Series]:
    labels = df[label_col].astype(str).str.strip()
    mapping = {label: i for i, label in enumerate(LABEL_ORDER)}
    encoded = labels.map(mapping)
    return labels, encoded


def validate_features(df: pd.DataFrame, requested_features: Sequence[str]) -> Tuple[List[str], Dict[str, str], Dict[str, int]]:
    usable: List[str] = []
    skipped: Dict[str, str] = {}
    valid_counts: Dict[str, int] = {}
    for feature in requested_features:
        if feature in EXCLUDED_FEATURES:
            skipped[feature] = "排除字段，不作为互信息输入"
            continue
        if feature not in df.columns:
            skipped[feature] = "字段不存在"
            print(f"跳过 {feature}: 字段不存在")
            continue
        numeric = pd.to_numeric(df[feature], errors="coerce").replace([np.inf, -np.inf], np.nan)
        valid_count = int(numeric.notna().sum())
        valid_counts[feature] = valid_count
        if valid_count == 0:
            skipped[feature] = "全为空或无法转为数值"
            print(f"跳过 {feature}: 全为空或无法转为数值")
        elif valid_count < MIN_VALID_COUNT:
            skipped[feature] = f"有效样本太少，仅 {valid_count} 条"
            print(f"跳过 {feature}: 有效样本太少，仅 {valid_count} 条")
        else:
            usable.append(feature)
    return usable, skipped, valid_counts


def compute_mi_for_features(
    df: pd.DataFrame,
    features: Sequence[str],
    encoded_label: pd.Series,
) -> Tuple[pd.DataFrame, int]:
    x = pd.DataFrame(index=df.index)
    for feature in features:
        x[feature] = pd.to_numeric(df[feature], errors="coerce").replace([np.inf, -np.inf], np.nan)

    analysis = x.copy()
    analysis["_label"] = encoded_label
    analysis = analysis.dropna(axis=0, how="any")
    if len(analysis) < MIN_VALID_COUNT:
        raise ValueError(f"互信息有效样本不足：{len(analysis)} 条。")
    if analysis["_label"].nunique() < 2:
        raise ValueError("互信息计算至少需要两个标签类别。")

    x_clean = analysis[list(features)]
    y_clean = analysis["_label"].astype(int)
    scores = mutual_info_classif(x_clean, y_clean, random_state=42, discrete_features=False)
    result = pd.DataFrame({"feature": list(features), "mutual_information": scores})
    result = result.sort_values("mutual_information", ascending=False).reset_index(drop=True)
    result["rank"] = np.arange(1, len(result) + 1)
    return result, len(analysis)


def save_mi_plot(table: pd.DataFrame, title: str, output_path: Path, figsize: Tuple[float, float]) -> None:
    if table.empty:
        return
    data = table.iloc[::-1].copy()
    fig, ax = plt.subplots(figsize=figsize, dpi=300)
    ax.barh(data["feature"], data["mutual_information"], color="#4E79A7", edgecolor="black", linewidth=0.4)
    ax.set_xlabel("Mutual Information")
    ax.set_ylabel("Feature")
    ax.set_title(title, pad=12)
    ax.grid(axis="x", linestyle="--", alpha=0.25)

    max_score = float(data["mutual_information"].max()) if len(data) else 0.0
    offset = max(max_score * 0.01, 0.005)
    for y_pos, value in enumerate(data["mutual_information"]):
        ax.text(value + offset, y_pos, f"{value:.3f}", va="center", fontsize=8)

    ax.set_xlim(0, max_score + max(max_score * 0.16, 0.08))
    fig.tight_layout()
    fig.savefig(output_path, bbox_inches="tight")
    plt.close(fig)


def markdown_table(df: pd.DataFrame, max_rows: int = 10) -> str:
    if df.empty:
        return "无。"
    show = df.head(max_rows).copy()
    for col in show.columns:
        if pd.api.types.is_float_dtype(show[col]):
            show[col] = show[col].map(lambda x: f"{x:.6f}")
        else:
            show[col] = show[col].map(str)
    lines = [
        "| " + " | ".join(show.columns) + " |",
        "| " + " | ".join("---" for _ in show.columns) + " |",
    ]
    for _, row in show.iterrows():
        lines.append("| " + " | ".join(str(row[col]) for col in show.columns) + " |")
    return "\n".join(lines)


def skipped_table(skipped: Dict[str, str]) -> str:
    if not skipped:
        return "无。"
    lines = ["| 字段 | 跳过原因 |", "| --- | --- |"]
    for field, reason in skipped.items():
        lines.append(f"| `{field}` | {reason} |")
    return "\n".join(lines)


def write_summary(
    output_dir: Path,
    actual_input: Path,
    label_col: str,
    ratio_note: str,
    candidate_features: Sequence[str],
    selected_features: Sequence[str],
    skipped_candidate: Dict[str, str],
    skipped_selected: Dict[str, str],
    candidate_result: pd.DataFrame,
    selected_result: pd.DataFrame,
    candidate_n: int,
    selected_n: int,
    font_name: str,
) -> None:
    def bullet_list(items: Sequence[str]) -> str:
        return "\n".join(f"- `{item}`" for item in items) if items else "- 无"

    report = f"""# 互信息特征重要性分析汇总

## 数据与标签

- 输入文件：`{actual_input}`
- 标签列：`{label_col}`
- 标签顺序：`{"、".join(LABEL_ORDER)}`
- 字体：`{font_name}`
- `RD_RS_ratio` 处理：{ratio_note}

## 候选特征集合实际参与计算字段

有效样本数：{candidate_n}

{bullet_list(candidate_features)}

## 第一版输入特征集合实际参与计算字段

有效样本数：{selected_n}

{bullet_list(selected_features)}

## 候选特征跳过字段

{skipped_table(skipped_candidate)}

## 第一版输入特征跳过字段

{skipped_table(skipped_selected)}

## 候选特征互信息 Top 10

{markdown_table(candidate_result[["rank", "feature", "mutual_information"]], 10)}

## 第一版输入特征互信息 Top 10

{markdown_table(selected_result[["rank", "feature", "mutual_information"]], 10)}

## 输出文件

- 候选特征 CSV：`{output_dir / "candidate_features_mi.csv"}`
- 候选特征 PNG：`{output_dir / "candidate_features_mi.png"}`
- 第一版输入特征 CSV：`{output_dir / "selected_v1_features_mi.csv"}`
- 第一版输入特征 PNG：`{output_dir / "selected_v1_features_mi.png"}`
"""
    (output_dir / "mi_summary.md").write_text(report, encoding="utf-8")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Compute mutual information rankings for logging features.")
    parser.add_argument("--input", default="outputs_precheck/modeling_data_point.csv", help="Preferred CSV or Excel input path.")
    parser.add_argument("--output", default="outputs_precheck/mi_analysis", help="Output directory.")
    parser.add_argument("--sheet", default="Sheet1", help="Excel sheet name used only for Excel input or fallback.")
    return parser.parse_args()


def main() -> None:
    warnings.filterwarnings("ignore", category=FutureWarning)
    configure_stdout()
    font_name = configure_matplotlib_fonts()
    args = parse_args()

    input_path = Path(args.input)
    output_dir = Path(args.output)
    output_dir.mkdir(parents=True, exist_ok=True)

    df, actual_input = load_data(input_path, args.sheet)
    df = clean_bad_values(df)
    df, ratio_note = ensure_ratio(df)

    label_col = find_label_column(df)
    if not label_col:
        raise ValueError("未找到标签列：y_fine 或 试油结论。")

    labels, encoded_label = encode_label(df, label_col)
    invalid_label_count = int(encoded_label.isna().sum())
    if invalid_label_count:
        print(f"提示：有 {invalid_label_count} 条标签不在固定六分类中，互信息计算时会排除。")

    candidate_features, skipped_candidate, _ = validate_features(df, CANDIDATE_FEATURES)
    selected_features, skipped_selected, _ = validate_features(df, SELECTED_V1_FEATURES)

    candidate_result, candidate_n = compute_mi_for_features(df, candidate_features, encoded_label)
    selected_result, selected_n = compute_mi_for_features(df, selected_features, encoded_label)

    candidate_csv = output_dir / "candidate_features_mi.csv"
    selected_csv = output_dir / "selected_v1_features_mi.csv"
    candidate_png = output_dir / "candidate_features_mi.png"
    selected_png = output_dir / "selected_v1_features_mi.png"

    candidate_result.to_csv(candidate_csv, index=False, encoding="utf-8-sig")
    selected_result.to_csv(selected_csv, index=False, encoding="utf-8-sig")
    save_mi_plot(
        candidate_result,
        "Mutual Information Ranking of Candidate Features",
        candidate_png,
        figsize=(10, max(8, 0.32 * len(candidate_result) + 2)),
    )
    save_mi_plot(
        selected_result,
        "Mutual Information Ranking of Selected V1 Features",
        selected_png,
        figsize=(8, 6),
    )

    write_summary(
        output_dir=output_dir,
        actual_input=actual_input,
        label_col=label_col,
        ratio_note=ratio_note,
        candidate_features=candidate_features,
        selected_features=selected_features,
        skipped_candidate=skipped_candidate,
        skipped_selected=skipped_selected,
        candidate_result=candidate_result,
        selected_result=selected_result,
        candidate_n=candidate_n,
        selected_n=selected_n,
        font_name=font_name,
    )

    print(f"读取文件：{actual_input}")
    print(f"标签列：{label_col}")
    print(ratio_note)
    print(f"候选特征参与分析数量：{len(candidate_features)}，有效样本数：{candidate_n}")
    print(f"第一版输入特征参与分析数量：{len(selected_features)}，有效样本数：{selected_n}")
    print(f"候选特征互信息图：{candidate_png.resolve()}")
    print(f"第一版输入特征互信息图：{selected_png.resolve()}")
    print(f"汇总说明：{(output_dir / 'mi_summary.md').resolve()}")


if __name__ == "__main__":
    main()
