# -*- coding: utf-8 -*-
"""Plot a compact Pearson correlation heatmap for selected logging features.

Example:
    python plot_selected_corr.py --input outputs_precheck/modeling_data_point.csv --output outputs_precheck
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path
from typing import List, Tuple

import matplotlib

matplotlib.use("Agg")

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd


SELECTED_FEATURES = [
    "GR",
    "SP",
    "AC",
    "CAL",
    "DEN",
    "CNL",
    "PE",
    "RD",
    "RS",
    "SW",
    "RD_RS_ratio",
]


def configure_stdout() -> None:
    try:
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    except Exception:
        pass


def configure_matplotlib_fonts() -> str:
    """Try to use a Chinese-capable font without failing if none is installed."""
    candidates = [
        "Microsoft YaHei",
        "SimHei",
        "SimSun",
        "Noto Sans CJK SC",
        "Source Han Sans SC",
        "Arial Unicode MS",
        "DejaVu Sans",
    ]
    available = {f.name for f in matplotlib.font_manager.fontManager.ttflist}
    selected = next((font for font in candidates if font in available), "DejaVu Sans")
    plt.rcParams["font.sans-serif"] = [selected]
    plt.rcParams["axes.unicode_minus"] = False
    return selected


def read_excel_fallback(path: Path, sheet_name: str) -> pd.DataFrame:
    """Read Excel with pandas first, then fall back to openpyxl for older envs."""
    try:
        return pd.read_excel(path, sheet_name=sheet_name)
    except Exception:
        import openpyxl

        wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
        if sheet_name not in wb.sheetnames:
            wb.close()
            raise ValueError(f"Sheet '{sheet_name}' not found in {path}.")
        ws = wb[sheet_name]
        rows = ws.iter_rows()
        header = [str(cell.value).strip() if cell.value is not None else f"Unnamed_{i+1}" for i, cell in enumerate(next(rows))]
        records = []
        for row in rows:
            values = [cell.value for cell in row[: len(header)]]
            if len(values) < len(header):
                values.extend([np.nan] * (len(header) - len(values)))
            records.append(values)
        wb.close()
        return pd.DataFrame(records, columns=header)


def load_data(input_path: Path, sheet_name: str) -> Tuple[pd.DataFrame, Path]:
    """Load preferred CSV input; if absent, fall back to the original workbook."""
    if input_path.exists():
        if input_path.suffix.lower() in [".csv", ".txt"]:
            return pd.read_csv(input_path), input_path
        if input_path.suffix.lower() in [".xlsx", ".xlsm", ".xls"]:
            return read_excel_fallback(input_path, sheet_name), input_path
        raise ValueError(f"Unsupported input type: {input_path.suffix}")

    fallback = Path("流体识别数据.xlsx")
    if fallback.exists():
        print(f"提示：未找到 {input_path}，改用原始 Excel：{fallback}")
        return read_excel_fallback(fallback, sheet_name), fallback
    raise FileNotFoundError(f"Input file not found: {input_path}; fallback Excel also not found: {fallback}")


def is_bad_value(value: object) -> bool:
    if value is None:
        return True
    if isinstance(value, str):
        text = value.strip()
        if text == "":
            return True
        if "#REF!" in text.upper():
            return True
        try:
            return float(text) == -9999.0
        except ValueError:
            return False
    try:
        return float(value) == -9999.0
    except (TypeError, ValueError):
        return False


def prepare_selected_data(df: pd.DataFrame) -> Tuple[pd.DataFrame, List[str], List[str]]:
    """Keep selected fields, coerce to numeric, and compute RD_RS_ratio if needed."""
    working = df.copy()

    for col in working.columns:
        working[col] = working[col].map(lambda x: np.nan if is_bad_value(x) else x)

    if "RD_RS_ratio" not in working.columns and {"RD", "RS"}.issubset(working.columns):
        rd = pd.to_numeric(working["RD"], errors="coerce")
        rs = pd.to_numeric(working["RS"], errors="coerce").replace(0, np.nan)
        working["RD_RS_ratio"] = rd / rs
        print("提示：输入数据不存在 RD_RS_ratio，已根据 RD / RS 自动计算。")

    missing = [col for col in SELECTED_FEATURES if col not in working.columns]
    present = [col for col in SELECTED_FEATURES if col in working.columns]
    selected = pd.DataFrame(index=working.index)
    for col in present:
        selected[col] = pd.to_numeric(working[col], errors="coerce")
        selected[col] = selected[col].replace([np.inf, -np.inf], np.nan)

    return selected, present, missing


def plot_heatmap(corr: pd.DataFrame, output_path: Path) -> None:
    fig, ax = plt.subplots(figsize=(10, 8), dpi=300)
    image = ax.imshow(corr.values, cmap="coolwarm", vmin=-1, vmax=1)

    ax.set_xticks(np.arange(len(corr.columns)))
    ax.set_yticks(np.arange(len(corr.index)))
    ax.set_xticklabels(corr.columns, rotation=45, ha="right", rotation_mode="anchor")
    ax.set_yticklabels(corr.index)
    ax.set_title("Pearson 相关性矩阵（精简输入特征）", pad=14)

    ax.set_xticks(np.arange(-0.5, len(corr.columns), 1), minor=True)
    ax.set_yticks(np.arange(-0.5, len(corr.index), 1), minor=True)
    ax.grid(which="minor", color="white", linestyle="-", linewidth=0.8)
    ax.tick_params(which="minor", bottom=False, left=False)

    for i in range(len(corr.index)):
        for j in range(len(corr.columns)):
            value = corr.iloc[i, j]
            label = "" if pd.isna(value) else f"{value:.2f}"
            color = "white" if pd.notna(value) and abs(value) >= 0.55 else "black"
            ax.text(j, i, label, ha="center", va="center", color=color, fontsize=8)

    cbar = fig.colorbar(image, ax=ax, fraction=0.046, pad=0.04)
    cbar.set_label("Pearson r")
    fig.tight_layout()
    fig.savefig(output_path, bbox_inches="tight")
    plt.close(fig)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Plot selected-feature Pearson correlation matrix.")
    parser.add_argument("--input", default="outputs_precheck/modeling_data_point.csv", help="Preferred CSV or Excel input path.")
    parser.add_argument("--output", default="outputs_precheck", help="Directory to write CSV and PNG outputs.")
    parser.add_argument("--sheet", default="Sheet1", help="Excel sheet name used only for Excel input or fallback.")
    return parser.parse_args()


def main() -> None:
    configure_stdout()
    font_name = configure_matplotlib_fonts()
    args = parse_args()

    input_path = Path(args.input)
    output_dir = Path(args.output)
    output_dir.mkdir(parents=True, exist_ok=True)

    df, actual_input = load_data(input_path, args.sheet)
    selected_df, present, missing = prepare_selected_data(df)

    if missing:
        print("提示：以下字段缺失，未参与计算：" + "、".join(missing))
    if selected_df.empty:
        raise ValueError("没有可用于相关性计算的字段。")

    nan_summary = selected_df.isna().sum().rename("NaN数量").reset_index().rename(columns={"index": "字段"})
    many_nan = nan_summary[nan_summary["NaN数量"] > len(selected_df) * 0.3]
    if not many_nan.empty:
        print("提示：以下字段 NaN 比例超过 30%：")
        for _, row in many_nan.iterrows():
            print(f"  - {row['字段']}: {int(row['NaN数量'])}/{len(selected_df)}")

    corr = selected_df.corr(method="pearson")
    csv_path = output_dir / "correlation_matrix_selected_features.csv"
    png_path = output_dir / "correlation_matrix_selected_features.png"
    corr.to_csv(csv_path, encoding="utf-8-sig")
    plot_heatmap(corr, png_path)

    print(f"读取文件：{actual_input}")
    print(f"中文字体：{font_name}")
    print("实际参与计算字段：" + "、".join(present))
    print(f"输出 CSV：{csv_path.resolve()}")
    print(f"输出 PNG：{png_path.resolve()}")
    print("各字段 NaN 数量：")
    for _, row in nan_summary.iterrows():
        print(f"  - {row['字段']}: {int(row['NaN数量'])}/{len(selected_df)}")


if __name__ == "__main__":
    main()
